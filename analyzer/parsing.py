"""Resume text extraction for ResumeLens.

Extract raw text from an uploaded resume file. Three formats are supported:

* PDF  -- via :mod:`pdfplumber`
* DOCX -- via :mod:`python-docx`
* TXT  -- decoded as UTF-8 (with a lenient fallback)

Requirement 1.2: if the file type is unsupported, or no text can be
extracted, the system must show a clear error and create nothing. This module
raises :class:`UnsupportedFileType` or :class:`NoTextExtracted` -- both
subclasses of :class:`ParsingError` -- so callers can surface a clear message
and skip creating any records.
"""

from __future__ import annotations

import io
import os
from typing import IO

import docx
import pdfplumber

# File extensions this module can parse, lower-cased and including the dot.
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")


class ParsingError(Exception):
    """Base class for recoverable resume-parsing failures (Requirement 1.2)."""


class UnsupportedFileType(ParsingError):
    """Raised when the file extension is not one we can parse."""


class NoTextExtracted(ParsingError):
    """Raised when a supported file yields no usable text."""


def _ext(filename: str) -> str:
    """Return the lower-cased file extension (including the dot)."""
    return os.path.splitext(filename or "")[1].lower()


def _as_bytes(file_obj: IO[bytes] | bytes) -> bytes:
    """Read ``file_obj`` into bytes, rewinding first when possible."""
    if isinstance(file_obj, (bytes, bytearray)):
        return bytes(file_obj)
    try:
        file_obj.seek(0)
    except (AttributeError, OSError):
        pass
    return file_obj.read()


def _extract_pdf(data: bytes) -> str:
    """Extract text from PDF bytes using pdfplumber."""
    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    """Extract text from DOCX bytes using python-docx (paragraphs + tables)."""
    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    """Decode TXT bytes as UTF-8, falling back to a lenient decode."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("utf-8", errors="replace")


def extract_text(file_obj: IO[bytes] | bytes, filename: str) -> str:
    """Return the raw text extracted from an uploaded resume.

    ``file_obj`` is the uploaded file (a Django ``UploadedFile``, any binary
    file-like object, or raw ``bytes``). ``filename`` is used to pick the
    extractor by extension.

    Raises :class:`UnsupportedFileType` for extensions outside
    :data:`SUPPORTED_EXTENSIONS`, and :class:`NoTextExtracted` when a supported
    file contains no usable text (Requirement 1.2).
    """
    ext = _ext(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. "
            "Please upload a PDF, DOCX or TXT resume."
        )

    data = _as_bytes(file_obj)

    try:
        if ext == ".pdf":
            text = _extract_pdf(data)
        elif ext == ".docx":
            text = _extract_docx(data)
        else:  # ".txt"
            text = _extract_txt(data)
    except (UnsupportedFileType, NoTextExtracted):
        raise
    except Exception as exc:  # malformed/corrupt file -> clear error
        raise NoTextExtracted(
            f"Could not read the uploaded {ext} file. It may be corrupt or "
            "password protected."
        ) from exc

    text = (text or "").strip()
    if not text:
        raise NoTextExtracted(
            "No text could be extracted from the uploaded file. "
            "If it is a scanned PDF, please upload a text-based resume."
        )
    return text
